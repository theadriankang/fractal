"""Chat attachments: upload, extraction and how files reach each model. Runs offline."""

import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import files
from app.main import app
from app.routers.chat import to_model_messages
from app.schemas import ChatTurn

client = TestClient(app)


def pdf_with_text(text: str) -> bytes:
    """A minimal one-page PDF whose content stream draws `text`."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n%s\nendobj\n" % (i, body))
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objects) + 1, xref))
    return out.getvalue()


def docx_with(paragraph: str) -> bytes:
    doc = Document()
    doc.add_paragraph(paragraph)
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Chiller 2", "Tripped"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture(autouse=True)
def tmp_uploads(tmp_path, monkeypatch):
    monkeypatch.setattr(files, "UPLOAD_DIR", tmp_path)


def upload(name: str, data: bytes):
    return client.post("/api/files", files={"file": (name, data)})


def test_pdf_upload_extracts_text_and_pages():
    res = upload("sop.pdf", pdf_with_text("Reset only after the oil heater runs 30 min"))
    assert res.status_code == 200
    meta = res.json()
    assert meta["kind"] == "pdf" and meta["pages"] == 1 and meta["chars"] > 0
    assert "oil heater" in files.text_attachments([meta["id"]])


def test_docx_upload_includes_paragraphs_and_tables():
    meta = upload("handover.docx", docx_with("Level 23 warm since 2pm")).json()
    text = files.text_attachments([meta["id"]])
    assert "Level 23 warm since 2pm" in text and "Chiller 2 | Tripped" in text


@pytest.mark.parametrize(
    ("name", "data", "status"),
    [
        ("setup.exe", b"MZ", 400),
        ("empty.txt", b"", 400),
        ("big.txt", b"x" * (files.MAX_BYTES + 1), 413),
        ("broken.pdf", b"not a pdf", 400),
        ("headerless.pdf", pdf_with_text("hi")[len(b"%PDF-1.4\n"):], 400),
    ],
)
def test_rejected_uploads(name, data, status):
    res = upload(name, data)
    assert res.status_code == status
    assert name in res.json()["detail"]


def test_claude_gets_native_blocks_before_the_question():
    pdf = upload("sop.pdf", pdf_with_text("hello")).json()
    png = upload("panel.png", b"\x89PNG fake image bytes").json()
    note = upload("notes.txt", b"CHWST 8.1 C").json()
    turns = [ChatTurn(role="user", content="What should I check?", files=[pdf["id"], png["id"], note["id"]])]
    [msg] = to_model_messages(turns, native_files=True)
    types = [b["type"] for b in msg["content"]]
    assert types == ["document", "text", "image", "document", "text"]
    assert msg["content"][0]["source"]["media_type"] == "application/pdf"
    assert msg["content"][3]["source"] == {"type": "text", "media_type": "text/plain", "data": "CHWST 8.1 C"}
    assert msg["content"][-1] == {"type": "text", "text": "What should I check?"}


def test_other_models_get_extracted_text_and_an_image_note():
    note = upload("notes.txt", b"CHWST 8.1 C").json()
    png = upload("panel.png", b"\x89PNG fake").json()
    [msg] = to_model_messages([ChatTurn(role="user", content="Q?", files=[note["id"], png["id"]])], native_files=False)
    assert isinstance(msg["content"], str)
    assert '<attached_file name="notes.txt">\nCHWST 8.1 C\n</attached_file>' in msg["content"]
    assert "cannot view images" in msg["content"]
    assert msg["content"].endswith("Q?")


def test_long_text_is_truncated_with_a_note():
    meta = upload("log.csv", b"a" * (files.OTHER_TEXT_LIMIT + 10)).json()
    text = files.text_attachments([meta["id"]])
    assert f"first {files.OTHER_TEXT_LIMIT:,} of {files.OTHER_TEXT_LIMIT + 10:,} characters" in text


def test_files_on_earlier_turns_are_resent_and_unknown_ids_are_safe():
    note = upload("notes.txt", b"CHWST 8.1 C").json()
    turns = [
        ChatTurn(role="user", content="First", files=[note["id"]]),
        ChatTurn(role="assistant", content="Answer"),
        ChatTurn(role="user", content="Follow-up", files=["../../etc/passwd", "0" * 32]),
    ]
    first, _, last = to_model_messages(turns, native_files=False)
    assert "CHWST 8.1 C" in first["content"]
    assert last["content"].count(files.MISSING) == 2


def test_turns_without_files_are_unchanged():
    turns = [ChatTurn(role="assistant", content="stray"), ChatTurn(role="user", content="Hi")]
    assert to_model_messages(turns, native_files=True) == [{"role": "user", "content": "Hi"}]
