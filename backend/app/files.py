"""Chat attachments: stored under backend/uploads/<id>/ and read back by id on every chat turn.

Each upload keeps the original bytes and its extracted text. Claude receives PDFs and
images natively (it reads scanned pages and figures); other models receive the extracted
text, truncated, as Prompt 9 in docs/BUILD_PROMPTS.md describes.
"""

import base64
import io
import json
import re
import uuid
from html import escape
from pathlib import Path

from .config import BACKEND_DIR

UPLOAD_DIR = BACKEND_DIR / "uploads"
MAX_BYTES = 10 * 1024 * 1024
MAX_FILES_PER_MESSAGE = 5
CLAUDE_TEXT_LIMIT = 100_000  # characters of a text/DOCX file sent to Claude
OTHER_TEXT_LIMIT = 20_000  # characters of any file sent to non-Claude models

# extension -> (kind, media type)
TYPES = {
    ".pdf": ("pdf", "application/pdf"),
    ".docx": ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ".txt": ("text", "text/plain"),
    ".md": ("text", "text/markdown"),
    ".csv": ("text", "text/csv"),
    ".png": ("image", "image/png"),
    ".jpg": ("image", "image/jpeg"),
    ".jpeg": ("image", "image/jpeg"),
}

_ID = re.compile(r"^[0-9a-f]{32}$")


class UploadError(ValueError):
    pass


def _extract(kind: str, data: bytes) -> tuple[str, int | None]:
    """Returns (text, page count). Images have no text."""
    if kind == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages).strip(), len(reader.pages)
    if kind == "docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            lines += [" | ".join(cell.text for cell in row.cells) for row in table.rows]
        return "\n".join(lines).strip(), None
    if kind == "text":
        return data.decode("utf-8-sig", errors="replace").strip(), None
    return "", None


def save_upload(filename: str, data: bytes) -> dict:
    name = Path(filename).name[:200] or "file"
    ext = Path(name).suffix.lower()
    if ext not in TYPES:
        raise UploadError(f"{name}: unsupported file type. Use PDF, DOCX, TXT, MD, CSV, PNG or JPG.")
    if not data:
        raise UploadError(f"{name} is empty.")
    if len(data) > MAX_BYTES:
        raise UploadError(f"{name} is larger than 10 MB.")

    kind, media_type = TYPES[ext]
    # pypdf tolerates a missing header, but Claude rejects such files on every later turn.
    if kind == "pdf" and not data.lstrip()[:5] == b"%PDF-":
        raise UploadError(f"{name} is not a valid PDF.")
    try:
        text, pages = _extract(kind, data)
    except Exception as e:  # corrupt or password-protected files
        raise UploadError(f"Could not read {name}: {e}") from e

    file_id = uuid.uuid4().hex
    folder = UPLOAD_DIR / file_id
    folder.mkdir(parents=True)
    (folder / "original").write_bytes(data)
    (folder / "text.txt").write_text(text, encoding="utf-8")
    meta = {"id": file_id, "name": name, "size": len(data), "kind": kind, "mediaType": media_type, "chars": len(text), "pages": pages}
    (folder / "meta.json").write_text(json.dumps(meta), encoding="utf-8")
    return meta


def _load(file_id: str) -> tuple[dict, Path] | None:
    if not _ID.match(file_id):
        return None
    folder = UPLOAD_DIR / file_id
    try:
        return json.loads((folder / "meta.json").read_text(encoding="utf-8")), folder
    except (OSError, ValueError):
        return None


def _clip(text: str, limit: int) -> str:
    if len(text) <= limit:
        return text
    return f"{text[:limit]}\n\n[Truncated: showing the first {limit:,} of {len(text):,} characters.]"


MISSING = "[An attached file is no longer available on the server.]"


def claude_blocks(file_ids: list[str]) -> list[dict]:
    """Content blocks for a Claude user turn; documents go before the question text."""
    blocks: list[dict] = []
    for file_id in file_ids:
        found = _load(file_id)
        if not found:
            blocks.append({"type": "text", "text": MISSING})
            continue
        meta, folder = found
        if meta["kind"] in ("pdf", "image"):
            data = base64.standard_b64encode((folder / "original").read_bytes()).decode()
            if meta["kind"] == "pdf":
                blocks.append({"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": data}, "title": meta["name"]})
            else:
                blocks.append({"type": "text", "text": f"Attached image: {meta['name']}"})
                blocks.append({"type": "image", "source": {"type": "base64", "media_type": meta["mediaType"], "data": data}})
            continue
        text = (folder / "text.txt").read_text(encoding="utf-8")
        if not text:
            blocks.append({"type": "text", "text": f"[Attached file {meta['name']} contains no readable text.]"})
            continue
        blocks.append({"type": "document", "source": {"type": "text", "media_type": "text/plain", "data": _clip(text, CLAUDE_TEXT_LIMIT)}, "title": meta["name"]})
    return blocks


def text_attachments(file_ids: list[str]) -> str:
    """Extracted text of the attachments, for models without native file input."""
    parts: list[str] = []
    for file_id in file_ids:
        found = _load(file_id)
        if not found:
            parts.append(MISSING)
            continue
        meta, folder = found
        name = escape(meta["name"], quote=True)
        if meta["kind"] == "image":
            parts.append(f"[Attached image {name}: this model cannot view images in Fractal yet.]")
            continue
        text = (folder / "text.txt").read_text(encoding="utf-8")
        if not text:
            parts.append(f"[Attached file {name} has no extractable text; it may be a scanned document.]")
            continue
        parts.append(f'<attached_file name="{name}">\n{_clip(text, OTHER_TEXT_LIMIT)}\n</attached_file>')
    return "\n\n".join(parts)
